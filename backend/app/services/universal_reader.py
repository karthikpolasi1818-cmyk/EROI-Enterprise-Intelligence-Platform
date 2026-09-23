from __future__ import annotations

from io import BytesIO
from pathlib import Path
import json
import mimetypes

import pandas as pd


SUPPORTED_DATA_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xls",
    ".json",
    ".parquet",
}

SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
    ".log",
    ".xml",
    ".html",
}

SUPPORTED_EXTENSIONS = (
    SUPPORTED_DATA_EXTENSIONS
    | SUPPORTED_DOCUMENT_EXTENSIONS
)


def get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def detect_file_type(filename: str) -> str:
    extension = get_extension(filename)

    mapping = {
        ".csv": "CSV",
        ".xlsx": "Excel",
        ".xls": "Excel",
        ".json": "JSON",
        ".parquet": "Parquet",
        ".pdf": "PDF",
        ".docx": "Word",
        ".pptx": "PowerPoint",
        ".txt": "Text",
        ".md": "Markdown",
        ".log": "Log",
        ".xml": "XML",
        ".html": "HTML",
    }

    return mapping.get(extension, "Unknown")


def detect_mime_type(filename: str) -> str:
    mime_type, _ = mimetypes.guess_type(filename)
    return mime_type or "application/octet-stream"


def read_csv(contents: bytes) -> pd.DataFrame:
    return pd.read_csv(BytesIO(contents))


def read_excel(contents: bytes) -> pd.DataFrame:
    return pd.read_excel(BytesIO(contents))


def read_json(contents: bytes) -> pd.DataFrame:
    text = contents.decode("utf-8", errors="replace")
    data = json.loads(text)

    if isinstance(data, list):
        return pd.json_normalize(data)

    if isinstance(data, dict):
        for key in ["data", "records", "results", "items"]:
            if key in data and isinstance(data[key], list):
                return pd.json_normalize(data[key])

        return pd.json_normalize(data)

    raise ValueError("Unsupported JSON structure.")


def read_parquet(contents: bytes) -> pd.DataFrame:
    return pd.read_parquet(BytesIO(contents))


def read_pdf(contents: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(contents))

    pages = []
    full_text = []

    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        pages.append(
            {
                "page": index,
                "characters": len(text),
                "words": len(text.split()),
                "text": text,
            }
        )

        full_text.append(text)

    combined = "\n".join(full_text)

    return {
        "document_type": "PDF",
        "pages": len(reader.pages),
        "characters": len(combined),
        "words": len(combined.split()),
        "text": combined,
        "page_details": pages,
    }


def read_docx(contents: bytes) -> dict:
    from docx import Document

    document = Document(BytesIO(contents))

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    tables = []

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        rows = []

        for row in table.rows:
            rows.append(
                [cell.text.strip() for cell in row.cells]
            )

        tables.append(
            {
                "table": table_index,
                "rows": rows,
            }
        )

    text = "\n".join(paragraphs)

    return {
        "document_type": "Word",
        "paragraphs": len(paragraphs),
        "tables": len(tables),
        "characters": len(text),
        "words": len(text.split()),
        "text": text,
        "table_details": tables,
    }


def read_txt(contents: bytes) -> dict:
    text = contents.decode(
        "utf-8",
        errors="replace",
    )

    return {
        "document_type": "Text",
        "lines": len(text.splitlines()),
        "characters": len(text),
        "words": len(text.split()),
        "text": text,
    }


def read_pptx(contents: bytes) -> dict:
    from pptx import Presentation

    presentation = Presentation(BytesIO(contents))

    slides = []
    all_text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)
                    all_text.append(text)

        combined = "\n".join(slide_text)

        slides.append(
            {
                "slide": slide_number,
                "text": combined,
                "characters": len(combined),
                "words": len(combined.split()),
            }
        )

    combined = "\n".join(all_text)

    return {
        "document_type": "PowerPoint",
        "slides": len(slides),
        "characters": len(combined),
        "words": len(combined.split()),
        "text": combined,
        "slide_details": slides,
    }


def read_plain_text(contents: bytes) -> dict:
    text = contents.decode(
        "utf-8",
        errors="replace",
    )

    return {
        "document_type": "Text",
        "characters": len(text),
        "words": len(text.split()),
        "lines": len(text.splitlines()),
        "text": text,
    }


def read_uploaded_file(
    filename: str,
    contents: bytes,
):
    extension = get_extension(filename)

    if extension == ".csv":
        return read_csv(contents)

    if extension in {".xlsx", ".xls"}:
        return read_excel(contents)

    if extension == ".json":
        return read_json(contents)

    if extension == ".parquet":
        return read_parquet(contents)

    if extension == ".pdf":
        return read_pdf(contents)

    if extension == ".docx":
        return read_docx(contents)

    if extension == ".pptx":
        return read_pptx(contents)

    if extension in {
        ".txt",
        ".md",
        ".log",
        ".xml",
        ".html",
    }:
        return read_plain_text(contents)

    raise ValueError(
        f"Unsupported file format: {extension}"
    )