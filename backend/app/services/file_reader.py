from io import BytesIO
from pathlib import Path
import json

import pandas as pd


SUPPORTED_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xls",
    ".pdf",
    ".docx",
    ".txt",
    ".json",
    ".pptx",
}


def get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def get_file_type(filename: str) -> str:
    extension = get_extension(filename)

    mapping = {
        ".csv": "CSV",
        ".xlsx": "Excel",
        ".xls": "Excel",
        ".pdf": "PDF",
        ".docx": "Word",
        ".txt": "Text",
        ".json": "JSON",
        ".pptx": "PowerPoint",
    }

    return mapping.get(extension, "Unknown")


def validate_extension(filename: str):
    extension = get_extension(filename)

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format: {extension}. "
            f"Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )


# ============================================================
# STRUCTURED FILE READERS
# ============================================================

def read_csv(contents: bytes) -> pd.DataFrame:
    return pd.read_csv(BytesIO(contents))


def read_excel(contents: bytes) -> pd.DataFrame:
    return pd.read_excel(BytesIO(contents))


def read_json(contents: bytes) -> pd.DataFrame:
    text = contents.decode("utf-8")

    data = json.loads(text)

    if isinstance(data, list):
        return pd.json_normalize(data)

    if isinstance(data, dict):

        # Try common record structures
        for key in ["data", "records", "results", "items"]:

            if key in data and isinstance(data[key], list):
                return pd.json_normalize(data[key])

        return pd.json_normalize(data)

    raise ValueError("Unsupported JSON structure.")


# ============================================================
# DOCUMENT READERS
# ============================================================

def read_pdf(contents: bytes) -> dict:
    from pypdf import PdfReader

    pdf = PdfReader(BytesIO(contents))

    pages = []
    full_text_parts = []

    for page_number, page in enumerate(pdf.pages, start=1):

        text = page.extract_text() or ""

        pages.append(
            {
                "page": page_number,
                "text": text,
                "characters": len(text),
                "words": len(text.split()),
            }
        )

        full_text_parts.append(text)

    full_text = "\n".join(full_text_parts)

    return {
        "document_type": "PDF",
        "pages": len(pdf.pages),
        "text": full_text,
        "characters": len(full_text),
        "words": len(full_text.split()),
        "page_details": pages,
    }


def read_docx(contents: bytes) -> dict:
    from docx import Document

    document = Document(BytesIO(contents))

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    tables = []

    for table_index, table in enumerate(document.tables, start=1):

        rows = []

        for row in table.rows:
            rows.append(
                [cell.text.strip() for cell in row.cells]
            )

        tables.append(
            {
                "table": table_index,
                "rows": rows,
            }
        )

    full_text = "\n".join(paragraphs)

    return {
        "document_type": "Word",
        "paragraphs": len(paragraphs),
        "tables": len(tables),
        "text": full_text,
        "characters": len(full_text),
        "words": len(full_text.split()),
        "table_details": tables,
    }


def read_txt(contents: bytes) -> dict:
    text = contents.decode(
        "utf-8",
        errors="replace",
    )

    lines = text.splitlines()

    return {
        "document_type": "Text",
        "lines": len(lines),
        "text": text,
        "characters": len(text),
        "words": len(text.split()),
    }


def read_pptx(contents: bytes) -> dict:
    from pptx import Presentation

    presentation = Presentation(BytesIO(contents))

    slides = []

    all_text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)
                    all_text.append(text)

        slides.append(
            {
                "slide": slide_number,
                "text": "\n".join(slide_text),
                "characters": len(
                    "\n".join(slide_text)
                ),
                "words": len(
                    "\n".join(slide_text).split()
                ),
            }
        )

    full_text = "\n".join(all_text)

    return {
        "document_type": "PowerPoint",
        "slides": len(slides),
        "text": full_text,
        "characters": len(full_text),
        "words": len(full_text.split()),
        "slide_details": slides,
    }


# ============================================================
# UNIVERSAL FILE READER
# ============================================================

def read_uploaded_file(
    filename: str,
    contents: bytes,
):
    """
    Universal EROI file reader.

    Returns:
        DataFrame for structured datasets
        dict for documents
    """

    validate_extension(filename)

    extension = get_extension(filename)

    if extension == ".csv":
        return read_csv(contents)

    if extension in {".xlsx", ".xls"}:
        return read_excel(contents)

    if extension == ".json":
        return read_json(contents)

    if extension == ".pdf":
        return read_pdf(contents)

    if extension == ".docx":
        return read_docx(contents)

    if extension == ".txt":
        return read_txt(contents)

    if extension == ".pptx":
        return read_pptx(contents)

    raise ValueError(
        f"No reader available for {extension}"
    )


# ============================================================
# NORMALIZED FILE PROFILE
# ============================================================

def profile_file(
    filename: str,
    contents: bytes,
) -> dict:

    file_type = get_file_type(filename)

    result = read_uploaded_file(
        filename,
        contents,
    )

    # --------------------------------------------------------
    # DATAFRAME
    # --------------------------------------------------------

    if isinstance(result, pd.DataFrame):

        df = result

        rows = len(df)
        columns = len(df.columns)

        total_cells = rows * columns

        missing_cells = int(
            df.isna().sum().sum()
        )

        duplicate_rows = int(
            df.duplicated().sum()
        )

        missing_percentage = (
            (missing_cells / total_cells) * 100
            if total_cells
            else 0
        )

        duplicate_percentage = (
            (duplicate_rows / rows) * 100
            if rows
            else 0
        )

        quality_score = max(
            0,
            100
            - missing_percentage
            - duplicate_percentage,
        )

        numeric_columns = (
            df.select_dtypes(
                include="number"
            ).columns.tolist()
        )

        numeric_summary = {}

        for column in numeric_columns:

            numeric_summary[column] = {
                "sum": float(
                    df[column].sum()
                ),
                "mean": float(
                    df[column].mean()
                )
                if len(df)
                else 0,
                "min": float(
                    df[column].min()
                )
                if len(df)
                else 0,
                "max": float(
                    df[column].max()
                )
                if len(df)
                else 0,
            }

        return {
            "success": True,
            "filename": filename,
            "file_type": file_type,
            "kind": "dataset",
            "dataset": {
                "rows": rows,
                "columns": columns,
                "column_names": [
                    str(column)
                    for column in df.columns
                ],
            },
            "quality": {
                "missing_cells": missing_cells,
                "missing_percentage": round(
                    missing_percentage,
                    2,
                ),
                "duplicate_rows": duplicate_rows,
                "duplicate_percentage": round(
                    duplicate_percentage,
                    2,
                ),
                "quality_score": round(
                    quality_score,
                    2,
                ),
            },
            "data_types": {
                str(column): str(dtype)
                for column, dtype in df.dtypes.items()
            },
            "numeric_summary": numeric_summary,
        }

    # --------------------------------------------------------
    # DOCUMENT
    # --------------------------------------------------------

    return {
        "success": True,
        "filename": filename,
        "file_type": file_type,
        "kind": "document",
        "document": result,
    }